#!/usr/bin/env python3
"""Export an evidence-linked document model into versioned DOCX/PDF drafts."""
import argparse
import json
from pathlib import Path
import sys
sys.dont_write_bytecode = True
import tempfile
from xml.sax.saxutils import escape
import workspace as ws

def items(model):
    yield model['name']; yield model['contact']
    for section in model['sections']:
        yield from section.get('paragraphs',[])
        yield from section.get('bullets',[])
        for entry in section.get('entries',[]):
            yield entry['title']
            if entry.get('subtitle'): yield entry['subtitle']
            yield from entry.get('paragraphs',[])
            yield from entry.get('bullets',[])

def validate_model(model,state):
    ws.require(isinstance(model,dict),'Document model must be an object.')
    ws.require(model.get('kind') in ('resume','cv','cover_letter'),'Invalid document kind.')
    ws.require(model.get('page_size') in ('Letter','A4'),'Choose Letter or A4.')
    ws.require(isinstance(model.get('sections'),list) and model['sections'],'Document sections required.')
    facts={f['id']:f for f in state['profile']['facts']}
    for section in model['sections']:
        ws.require(isinstance(section.get('heading'),str),'Section heading must be text.')
    for item in items(model):
        ws.require(isinstance(item,dict) and isinstance(item.get('text'),str) and item['text'].strip(),'Every content item needs text.')
        ids=item.get('fact_ids')
        ws.require(isinstance(ids,list) and ids,'Every candidate-facing item needs supporting fact IDs.')
        ws.require(all(x in facts and facts[x]['status'] in ('verified','candidate_reported') for x in ids),'Document uses an unknown, unresolved or superseded fact.')
    return model

def build_docx(model,path):
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d=Document(); section=d.sections[0]
    if model['page_size']=='A4': section.page_width=Inches(8.2677); section.page_height=Inches(11.6929)
    else: section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.7)
    normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
    for name in ('Title','Heading 1','List Bullet'):
        d.styles[name].font.name='Arial'
        pp=d.styles[name].element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
        if pp is not None:
            border=pp.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr')
            if border is not None: pp.remove(border)
    title=d.styles['Title']; title.font.size=Pt(14); title.font.bold=True; title.font.color.rgb=RGBColor(0,0,0)
    heading=d.styles['Heading 1']; heading.font.size=Pt(11); heading.font.bold=True; heading.font.color.rgb=RGBColor.from_string('183F5A')
    heading.paragraph_format.space_before=Pt(11); heading.paragraph_format.space_after=Pt(5); heading.paragraph_format.keep_with_next=True
    bullet=d.styles['List Bullet']; bullet.font.size=Pt(11); bullet.paragraph_format.space_after=Pt(3)
    p=d.add_paragraph(model['name']['text'],'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(model['contact']['text']); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for s in model['sections']:
        if s['heading']: d.add_paragraph(s['heading'],'Heading 1')
        for x in s.get('paragraphs',[]): d.add_paragraph(x['text'])
        for x in s.get('bullets',[]): d.add_paragraph(x['text'],'List Bullet')
        for e in s.get('entries',[]):
            p=d.add_paragraph(); p.add_run(e['title']['text']).bold=True; p.paragraph_format.keep_with_next=True
            if e.get('subtitle'):
                p=d.add_paragraph(e['subtitle']['text']); p.paragraph_format.keep_with_next=True
            for x in e.get('paragraphs',[]): d.add_paragraph(x['text'])
            for x in e.get('bullets',[]): d.add_paragraph(x['text'],'List Bullet')
    d.core_properties.author=''; d.core_properties.last_modified_by=''; d.core_properties.title=model['kind'].replace('_',' ').title()
    d.save(path)

def build_pdf(model,path,arial=None):
    from reportlab.lib.pagesizes import A4,letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,ListFlowable,ListItem
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    font='Helvetica'; bold='Helvetica-Bold'
    if arial:
        font='CandidateArial'; pdfmetrics.registerFont(TTFont(font,str(arial)))
        sibling=Path(arial).with_name('arialbd.ttf')
        if sibling.exists(): bold='CandidateArialBold'; pdfmetrics.registerFont(TTFont(bold,str(sibling)))
    normal=ParagraphStyle('Body',fontName=font,fontSize=11,leading=14,spaceAfter=5)
    title=ParagraphStyle('Name',parent=normal,fontName=bold,fontSize=14,leading=18,alignment=1,spaceAfter=5)
    contact=ParagraphStyle('Contact',parent=normal,alignment=1,spaceAfter=8)
    heading=ParagraphStyle('Heading',parent=normal,fontName=bold,textColor=HexColor('#183F5A'),spaceBefore=10,spaceAfter=5,keepWithNext=True)
    entry=ParagraphStyle('Entry',parent=normal,fontName=bold,keepWithNext=True)
    sub=ParagraphStyle('Subtitle',parent=normal,keepWithNext=True)
    bullet=ParagraphStyle('BulletBody',parent=normal,spaceAfter=3)
    def para(x,style=normal):return Paragraph(escape(x['text']).replace('\n','<br/>'),style)
    def bullets(values):
        if values: story.append(ListFlowable([ListItem(para(x,bullet),leftIndent=0) for x in values],bulletType='bullet',leftIndent=13,bulletFontName=font,bulletFontSize=8,spaceAfter=4))
    story=[para(model['name'],title),para(model['contact'],contact)]
    for s in model['sections']:
        if s['heading']: story.append(Paragraph(escape(s['heading']),heading))
        story.extend(para(x) for x in s.get('paragraphs',[])); bullets(s.get('bullets',[]))
        for e in s.get('entries',[]):
            story.append(para(e['title'],entry))
            if e.get('subtitle'): story.append(para(e['subtitle'],sub))
            story.extend(para(x) for x in e.get('paragraphs',[])); bullets(e.get('bullets',[]))
    SimpleDocTemplate(str(path),pagesize=letter if model['page_size']=='Letter' else A4,rightMargin=50.4,leftMargin=50.4,topMargin=46.8,bottomMargin=46.8,title=model['kind'].replace('_',' ').title(),author='').build(story)

def export(value,model,stem,arial=None):
    root=ws.checked_root(value); state=ws.load(root); validate_model(model,state)
    ws.require(stem.startswith(('applications/','documents/')),'Document stem must be inside applications/ or documents/.')
    for suffix in ('.docx','.pdf','.model.json','.claims.json'):
        ws.require(not ws.inside(root,stem+suffix).exists(),'Use a new version: an output filename already exists.')
    scratch=ws.inside(root,'scratch'); scratch.mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='export-',dir=scratch) as temp:
        temp=Path(temp); build_docx(model,temp/'draft.docx'); build_pdf(model,temp/'draft.pdf',arial)
        from pypdf import PdfReader
        reader=PdfReader(temp/'draft.pdf'); text='\n'.join(p.extract_text() or '' for p in reader.pages)
        normalized=' '.join(text.split())
        ws.require(all(' '.join(i['text'].split()) in normalized for i in items(model)),'PDF text extraction did not preserve all model content; inspect before use.')
        claims={'workspace_id':state['workspace_id'],'state_revision':state['revision'],'items':[{'text':i['text'],'fact_ids':i['fact_ids']} for i in items(model)],'pdf_pages':len(reader.pages),'text_extraction':'passed','visual_review':'not_performed','semantic_review':'not_performed','docx_pagination':'not_verified'}
        (temp/'draft.model.json').write_text(json.dumps(model,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        (temp/'draft.claims.json').write_text(json.dumps(claims,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        # No export may be committed after evidence changed while rendering.
        with ws.write_lock(root):
            ws.require(ws.load(root)['revision']==state['revision'],'State changed during export; reread facts and try a new export.')
            outputs=[ws.import_file(root,temp/('draft'+suffix),stem+suffix) for suffix in ('.docx','.pdf','.model.json','.claims.json')]
    return {'outputs':outputs,'pdf_pages':len(reader.pages),'visual_review':'not_performed','semantic_review':'not_performed','status':'draft'}

def main():
    p=argparse.ArgumentParser(description=__doc__); p.add_argument('--workspace',required=True); p.add_argument('--model',required=True); p.add_argument('--stem',required=True); p.add_argument('--arial')
    a=p.parse_args()
    try:
        result=export(a.workspace,json.loads(Path(a.model).read_text(encoding='utf-8')),a.stem,a.arial); print(json.dumps(result,indent=2)); return 0
    except ImportError as e:
        print(f'Document dependency missing: {e}. Use a workspace-local runtime with requirements-documents.txt.',file=sys.stderr); return 2
    except (ws.WorkspaceError,OSError,ValueError,KeyError,TypeError) as e:
        print(f'Document export error: {e}',file=sys.stderr); return 2

if __name__=='__main__': sys.exit(main())
