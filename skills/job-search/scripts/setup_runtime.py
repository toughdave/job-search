#!/usr/bin/env python3
"""Install and verify this workspace's private document runtime. Python 3.9+."""
import argparse
import hashlib
import json
import os
import shutil
import socket
import uuid
from contextlib import contextmanager
from pathlib import Path
import subprocess
import sys
import tempfile
import venv

import workspace as ws

SKILL = Path(__file__).resolve().parents[1]
REQUIREMENTS = SKILL / 'requirements-documents.txt'


def run(command, lock=None):
    env=os.environ.copy()
    # Keep corporate indexes, certificates and proxies; never redirect installs out of our venv.
    for key in ('PIP_TARGET','PIP_PREFIX','PIP_ROOT','PIP_USER','PIP_PYTHON'):
        env.pop(key,None)
    owner=None
    if lock:
        owner=json.loads(lock.read_text(encoding='utf-8'));owner.update(phase='spawning',child_pid=None);ws.atomic_json(lock,owner)
    child=subprocess.Popen([str(x) for x in command],stdout=subprocess.PIPE,stderr=subprocess.PIPE,
                           text=True,encoding='utf-8',errors='replace',env=env)
    if lock:
        owner.update(phase='child',child_pid=child.pid);ws.atomic_json(lock,owner)
    output,error=child.communicate()
    if lock:
        owner.update(phase='idle',child_pid=None);ws.atomic_json(lock,owner)
    # Do not echo pip output: a configured private index may contain credentials.
    ws.require(child.returncode==0,'Runtime command failed ('+Path(str(command[0])).name+', exit '+str(child.returncode)+'). Check interpreter health, network/mirror access and package compatibility. No credentials are included in this message.')
    return output


def probe(scratch):
    """Exercise actual installed imports, serialization, text extraction and rendering."""
    from importlib.metadata import version
    from zoneinfo import ZoneInfo
    import re
    import docx
    import pypdf
    import pypdfium2
    import documents
    versions = {}
    for line in REQUIREMENTS.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line or line.startswith('#'):
            continue
        match = re.fullmatch(r'([\w-]+)(==|>=)([\d.]+)', line)
        ws.require(match is not None, 'Unsupported dependency declaration: ' + line)
        name, operator, expected = match.groups()
        actual = version(name)
        numeric = lambda value: tuple(int(part) for part in value.split('.'))
        ws.require(actual == expected if operator == '==' else numeric(actual) >= numeric(expected),
                   'Dependency version differs from this skill: ' + name)
        versions[name] = actual
    ZoneInfo('America/Toronto')
    ZoneInfo('Asia/Kolkata')
    with tempfile.TemporaryDirectory(prefix='runtime-check-', dir=scratch) as temp:
        temp = Path(temp)
        model = json.loads((SKILL / 'assets/document-example.json').read_text(encoding='utf-8'))
        documents.build_docx(model, temp / 'check.docx')
        documents.build_pdf(model, temp / 'check.pdf')
        ws.require(model['name']['text'] in '\n'.join(p.text for p in docx.Document(temp / 'check.docx').paragraphs),
                   'DOCX readback failed.')
        reader = pypdf.PdfReader(temp / 'check.pdf')
        ws.require(model['name']['text'] in reader.pages[0].extract_text(), 'PDF readback failed.')
        with pypdfium2.PdfDocument(str(temp / 'check.pdf')) as pdf:
            page = pdf[0]
            bitmap = page.render(scale=1)
            picture = bitmap.to_pil()
            picture.save(temp / 'check.png')
            ws.require(picture.width > 0 and picture.height > 0, 'PDF rendering failed.')
            picture.close()
            bitmap.close()
            page.close()
    return {'packages': versions, 'checks': ['imports', 'timezone_data', 'docx_readback',
            'pdf_text', 'pdf_render'], 'docx_visual_review': 'requires_available_renderer',
            'web_and_scheduler': 'check_in_host', 'candidate_script_fonts': 'check_with_actual_text'}


def cache_base():
    override=os.environ.get('JOB_SEARCH_CACHE_DIR')
    if override:return Path(override).expanduser().resolve()
    if os.name=='nt':return (Path(os.environ.get('LOCALAPPDATA',Path.home()/'AppData/Local'))/'job-search').resolve()
    if sys.platform=='darwin':return (Path.home()/'Library/Caches/job-search').resolve()
    return (Path(os.environ.get('XDG_CACHE_HOME',Path.home()/'.cache'))/'job-search').resolve()


def runtime_home(state):
    base=cache_base();key=hashlib.sha256(state['workspace_id'].encode()).hexdigest()[:32]
    home=ws.inside(base,'workspaces/'+key)
    return home


def prepare_home(root,state,check_only):
    home=runtime_home(state)
    ws.require(not home.is_relative_to(root),'Runtime cache must be outside the candidate records folder.')
    for parent in [home,*home.parents]:
        ws.require(not (parent/'SKILL.md').exists() and not (parent/'skills/job-search/SKILL.md').exists(),'Runtime cache cannot be inside the public source or installed skill.')
    marker=ws.inside(home,'setup-owner.json')
    if marker.exists():
        owner=json.loads(marker.read_text(encoding='utf-8'))
        ws.require(owner.get('format')==2 and owner.get('workspace_id')==state['workspace_id'],'Foreign runtime cache; no files changed.')
    else:
        ws.require(not check_only,'Run setup first; no verified private cache exists.')
        ws.require(not home.exists() or not any(home.iterdir()),'Existing unowned cache folder; no files adopted.')
        home.mkdir(parents=True,exist_ok=True)
        ws.atomic_json(marker,{'format':2,'workspace_id':state['workspace_id']})
    if not check_only:
        for name in ('.gitignore','.ignore'):(home/name).write_text('*\n',encoding='utf-8')
    return home


def unlock_runtime(home,token):
    lock=ws.inside(home,'setup.lock');recovery=ws.inside(home,'unlock.lock')
    try:fd=os.open(recovery,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
    except FileExistsError:raise ws.WorkspaceError('Another runtime recovery is active; inspect unlock.lock.')
    try:
        os.close(fd);original=lock.read_bytes();owner=json.loads(original)
        ws.require(owner.get('token')==token,'Runtime lock token differs.')
        ws.require(owner.get('host')==socket.gethostname(),'Recover the runtime lock on its original host.')
        ws.require(not ws.process_alive(owner.get('pid')),'Runtime setup owner is still running.')
        ws.require(owner.get('phase')!='spawning','Child ownership was interrupted during launch; verify children manually before recovery.')
        if owner.get('child_pid'):ws.require(not ws.process_alive(owner['child_pid']),'Installer child is still running; wait for it before recovery.')
        ws.require(lock.read_bytes()==original,'Runtime lock changed during recovery.')
        lock.unlink();return {'unlocked':True}
    finally:recovery.unlink()


@contextmanager
def runtime_lock(home):
    lock=ws.inside(home,'setup.lock');token=uuid.uuid4().hex
    try:fd=os.open(lock,os.O_CREAT|os.O_EXCL|os.O_WRONLY,0o600)
    except FileExistsError:raise ws.WorkspaceError('Runtime setup lock exists at '+str(lock)+'. Inspect its owner; recover a stopped owner with --unlock-token. Never remove a live lock.')
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as stream:
            json.dump({'pid':os.getpid(),'host':socket.gethostname(),'token':token,'created_at':ws.now(),'phase':'idle','child_pid':None},stream)
        yield lock
    finally:
        if lock.exists() and json.loads(lock.read_text(encoding='utf-8')).get('token')==token:lock.unlink()


def interpreter_path(target):
    return target/('Scripts/python.exe' if os.name=='nt' else 'bin/python')


def healthy(interpreter):
    if not interpreter.is_file():return False
    try:
        data=json.loads(run([interpreter,'-I','-c','import sys,ssl,json; print(json.dumps({"version":list(sys.version_info[:2]),"prefix":sys.prefix,"base":sys.base_prefix}))']))
        return data['version']>=[3,9] and data['prefix']!=data['base']
    except (OSError,ValueError):return False


def preflight():
    ws.require(sys.version_info>=(3,9),'Python 3.9 or newer is required.')
    import ssl,ensurepip
    with tempfile.TemporaryDirectory(prefix='job-search-python-') as temp:
        target=Path(temp)/'venv';venv.EnvBuilder(with_pip=True,symlinks=False).create(target)
        ws.require(healthy(interpreter_path(target)),'This Python cannot create a usable isolated environment.')
        run([interpreter_path(target),'-I','-m','pip','--version'])
    return {'usable':True,'python':sys.executable,'base_python':getattr(sys,'_base_executable',sys.executable),'version':sys.version.split()[0]}


def setup(value, check_only=False, uv=None, unlock_token=None):
    ws.require(sys.version_info >= (3,9),'Python 3.9 or newer is required.')
    root=ws.checked_root(value);state=ws.load(root)
    for parent in [root,*root.parents]:
        ws.require(not (parent/'SKILL.md').exists() and not (parent/'skills/job-search/SKILL.md').exists(),'Runtime must be outside the installed skill and public source checkout.')
    metadata=ws.inside(root,'.runtime');record=ws.inside(root,'.runtime/runtime.json')
    prior=json.loads(record.read_text(encoding='utf-8')) if record.is_file() else {}
    repair_packages=prior.get('status')=='incomplete'
    ws.require(not prior.get('workspace_id') or prior['workspace_id']==state['workspace_id'],'Foreign runtime record.')
    try:home=prepare_home(root,state,check_only)
    except (ws.WorkspaceError,OSError,ValueError) as error:
        if check_only and record.is_file():
            prior.update(status='incomplete',last_error=str(error),failed_at=ws.now())
            ws.atomic_json(record,prior)
        raise
    if unlock_token:return unlock_runtime(home,unlock_token)
    target=ws.inside(home,'venv');interpreter=interpreter_path(target)
    fingerprint=hashlib.sha256(REQUIREMENTS.read_bytes().replace(b'\r\n',b'\n')).hexdigest()
    if check_only:ws.require(record.is_file(),'Run setup first; no runtime record.')
    else:
        metadata.mkdir(exist_ok=True)
        (metadata/'.ignore').write_text('*\n',encoding='utf-8')
    with runtime_lock(home) as lock:
        try:
            if check_only:
                ws.require(prior.get('requirements_sha256')==fingerprint,'Dependencies changed; run setup again.')
                ws.require(prior.get('status')=='runtime-ready','Previous runtime setup is incomplete; repair it.')
                ws.require(healthy(interpreter),'Private interpreter is broken; run setup with a healthy bootstrap Python.')
            else:
                prior.update(status='installing',workspace_id=state['workspace_id'],cache_path=str(home))
                ws.atomic_json(record,prior)
                # The venv may contain its executable but point to a deleted base interpreter.
                if not healthy(interpreter):
                    if target.exists():
                        ws.require(not Path(sys.executable).resolve().is_relative_to(target.resolve()),'Use a healthy bootstrap interpreter outside the broken venv before rebuilding it.')
                        archive=ws.inside(home,'venv-preserved-'+uuid.uuid4().hex)
                        ws.require(target.resolve().is_relative_to(home.resolve()) and archive.resolve().is_relative_to(home.resolve()),'Runtime archive path escaped cache.')
                        target.rename(archive);prior.setdefault('preserved_environments',[]).append(str(archive))
                    if uv:run([uv,'venv','--seed','--python',sys.executable,target],lock)
                    else:run([sys.executable,'-I','-m','venv','--copies',target],lock)
                    ws.require(healthy(interpreter),'Selected Python cannot create a usable environment. Run --preflight on another interpreter or install managed Python.')
                run([interpreter,'-I','-m','pip','install','--disable-pip-version-check','--only-binary=:all:',
                     '--no-input','--no-user','--prefix',target,'--cache-dir',home/'downloads',
                     *(['--force-reinstall'] if repair_packages else []),'-r',REQUIREMENTS],lock)
            result=json.loads(run([interpreter,'-E','-s','-B',Path(__file__).resolve(),'--probe',home],lock))
            result.update(status='runtime-ready',python=str(interpreter),bootstrap_python=getattr(sys,'_base_executable',sys.executable),
                          cache_path=str(home),workspace_id=state['workspace_id'],requirements_sha256=fingerprint,checked_at=ws.now())
            result['preserved_environments']=prior.get('preserved_environments',[])
            # Move only this helper's owned legacy package directory out of synced records.
            old=ws.inside(root,'.runtime/job-search-venv');old_marker=ws.inside(root,'.runtime/setup-owner.json')
            if old.exists() and old_marker.is_file() and json.loads(old_marker.read_text(encoding='utf-8')).get('workspace_id')==state['workspace_id']:
                result['legacy_cleanup_pending']=str(old)
                legacy_locked=ws.inside(root,'.runtime/setup.lock').exists()
                if legacy_locked:result['legacy_cleanup_reason']='Legacy setup lock exists; verify the old installer has stopped before migration.'
                if not check_only and not legacy_locked and not Path(sys.executable).resolve().is_relative_to(old.resolve()):
                    dest=ws.inside(home,'legacy-preserved-'+uuid.uuid4().hex)
                    ws.require(old.resolve().is_relative_to(root.resolve()) and dest.resolve().is_relative_to(home.resolve()),'Legacy move escapes its owned folders.')
                    shutil.move(str(old),str(dest));result['preserved_environments'].append(str(dest));result.pop('legacy_cleanup_pending')
            ws.atomic_json(record,result);return result
        except Exception as error:
            # Preserve discovery and recovery paths, including the last known interpreter.
            prior.update(status='incomplete',workspace_id=state['workspace_id'],cache_path=str(home),last_error=str(error),failed_at=ws.now())
            ws.atomic_json(record,prior);raise


def main():
    ws.configure_output()
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workspace');parser.add_argument('--check-only',action='store_true')
    parser.add_argument('--uv');parser.add_argument('--unlock-token');parser.add_argument('--preflight',action='store_true')
    parser.add_argument('--probe',help=argparse.SUPPRESS)
    args=parser.parse_args()
    try:
        result=preflight() if args.preflight else probe(args.probe) if args.probe else setup(args.workspace,args.check_only,args.uv,args.unlock_token)
        print(json.dumps(result,ensure_ascii=False,indent=2));return 0
    except (ws.WorkspaceError,OSError,ValueError,ImportError,AttributeError,subprocess.SubprocessError) as error:
        print('Runtime setup incomplete: '+str(error),file=sys.stderr);return 2


if __name__=='__main__':sys.exit(main())
